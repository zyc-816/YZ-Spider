import os
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from utils import logger


@dataclass
class ExamSubject:
    subject_code: str  # kskmdm 考试科目代码
    subject_name: str  # kskmmc 考试科目名称
    reference_notes: Optional[str] = None  # cksm 参考说明

    @classmethod
    def from_dict(cls, data: Optional[Dict[str, Any]]) -> "ExamSubject":
        if not data:
            return cls(subject_code="", subject_name="", reference_notes=None)
        return cls(
            subject_code=str(data.get("kskmdm") or ""),
            subject_name=str(data.get("kskmmc") or ""),
            reference_notes=data.get("cksm"),
        )


@dataclass
class ExamSubjectGroup:
    subject1: ExamSubject  # km1Vo 科目一（思想政治理论/综合）
    subject2: ExamSubject  # km2Vo 科目二（外国语）
    subject3: ExamSubject  # km3Vo 科目三（数学/业务课一）
    subject4: ExamSubject  # km4Vo 科目四（专业课）

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "ExamSubjectGroup":
        return cls(
            subject1=ExamSubject.from_dict(data.get("km1Vo")),
            subject2=ExamSubject.from_dict(data.get("km2Vo")),
            subject3=ExamSubject.from_dict(data.get("km3Vo")),
            subject4=ExamSubject.from_dict(data.get("km4Vo")),
        )


@dataclass
class SchoolDetail:
    # basic information
    school_code: str = ""  # dwdm 招生单位代码
    school_name: str = ""  # dwmc 招生单位名称
    province_code: str = ""  # szssm 所在省市代码
    province_name: str = ""  # szss 所在省市名称
    has_doctoral_point: int = 0  # bs 是否有博士点
    is_double_first_class: int = 0  # syl 是否“双一流”高校
    is_985: int = 0  # b985 是否“985”高校
    is_211: int = 0  # b211 是否“211”高校
    has_graduate_school: int = 0  # yjsy 是否建有研究生院
    is_self_scoring: int = 0  # zhx 是否34所自主划线

    # major and exam info
    exam_type_code: str = ""  # ksfsdm 考试方式代码
    exam_type_name: str = ""  # ksfsmc 考试方式名称
    college_code: str = ""  # yxsdm 院系所代码
    college_name: str = ""  # yxsmc 院系所名称
    discipline_category_code: str = ""  # mldm 学科门类代码
    discipline_category_name: str = ""  # mlmc 学科门类名称
    first_discipline_code: str = ""  # yjxkdm 一级学科代码
    first_discipline_name: str = ""  # yjxkmc 一级学科名称
    major_code: str = ""  # zydm 专业代码
    major_name: str = ""  # zymc 专业名称
    degree_type: str = ""  # xwlx 学位类型（专硕/学硕）
    direction_code: str = ""  # yjfxdm 研究方向代码
    direction_name: str = ""  # yjfxmc 研究方向名称
    advisor_name: str = ""  # zdjs 指导教师
    major_remarks: str = ""  # zybz 专业备注
    subject_groups: List[ExamSubjectGroup] = field(default_factory=list)  # kskmz 考试科目组列表

    # study type and plan
    study_mode: str = ""  # xxfs 学习方式（全日制/非全日制）
    total_planned_enrollment: Optional[int] = None  # nzsrs 总拟招生人数
    planned_enrollment_desc: Optional[str] = None  # nzsrsstr 拟招生说明文本
    exempt_student_count: Optional[int] = None  # ssjstmrs 接收推免人数
    is_veteran_plan: int = 0  # tydxs 退役大学生士兵计划
    is_special_public_plan: int = 0  # jsggjh 骨干定向专项计划

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "SchoolDetail":
        def to_int(val: Any, default: int = 0) -> int:
            if val is None or val == "":
                return default
            try:
                return int(val)
            except (ValueError, TypeError):
                return default

        raw_groups = data.get("kskmz") or []
        subject_groups = [
            ExamSubjectGroup.from_dict(group) 
            for group in raw_groups 
            if isinstance(group, dict)
        ]

        return cls(
            school_code=str(data.get("dwdm") or ""),
            school_name=str(data.get("dwmc") or ""),
            province_code=str(data.get("szssm") or ""),
            province_name=str(data.get("szss") or ""),
            has_doctoral_point=to_int(data.get("bs")),
            is_double_first_class=to_int(data.get("syl")),
            is_985=to_int(data.get("b985")),
            is_211=to_int(data.get("b211")),
            has_graduate_school=to_int(data.get("yjsy")),
            is_self_scoring=to_int(data.get("zhx")),

            exam_type_code=str(data.get("ksfsdm") or ""),
            exam_type_name=str(data.get("ksfsmc") or ""),
            college_code=str(data.get("yxsdm") or ""),
            college_name=str(data.get("yxsmc") or ""),
            discipline_category_code=str(data.get("mldm") or ""),
            discipline_category_name=str(data.get("mlmc") or ""),
            first_discipline_code=str(data.get("yjxkdm") or ""),
            first_discipline_name=str(data.get("yjxkmc") or ""),
            major_code=str(data.get("zydm") or ""),
            major_name=str(data.get("zymc") or ""),
            degree_type=str(data.get("xwlx") or ""),
            direction_code=str(data.get("yjfxdm") or ""),
            direction_name=str(data.get("yjfxmc") or ""),
            advisor_name=str(data.get("zdjs") or ""),
            major_remarks=str(data.get("zybz") or ""),
            subject_groups=subject_groups,

            study_mode=str(data.get("xxfs") or ""),
            total_planned_enrollment=data.get("nzsrs") if data.get("nzsrs") is not None else None,
            planned_enrollment_desc=data.get("nzsrsstr"),
            exempt_student_count=data.get("ssjstmrs") if data.get("ssjstmrs") is not None else None,
            is_veteran_plan=to_int(data.get("tydxs")),
            is_special_public_plan=to_int(data.get("jsggjh")),
        )


class ExcelPipeline:
    def __init__(self, data: List[Dict[str, Any]], major_code: str, major_name: str, output_dir: str = config.OUTPUT_DIR):
        self.data = list(data)
        self.major_code = str(major_code)
        self.major_name = str(major_name)
        self.output_dir = output_dir

    def _get_school_detail_list(self) -> List[SchoolDetail]:
        detail_list = [SchoolDetail.from_dict(_) for _ in self.data]
        return detail_list

    def _flatten_detail(self, item: SchoolDetail) -> Dict[str, Any]:
        sub1_list, sub2_list, sub3_list, sub4_list = [], [], [], []
        multi_group = len(item.subject_groups) > 1  # 是否存在多种选考方案

        for idx, group in enumerate(item.subject_groups, start=1):
            prefix = f"方案{idx}: " if multi_group else ""
            sub1_list.append(f"{prefix}({group.subject1.subject_code}){group.subject1.subject_name}".strip())
            sub2_list.append(f"{prefix}({group.subject2.subject_code}){group.subject2.subject_name}".strip())
            sub3_list.append(f"{prefix}({group.subject3.subject_code}){group.subject3.subject_name}".strip())
            sub4_list.append(f"{prefix}({group.subject4.subject_code}){group.subject4.subject_name}".strip())

        degree_map = {"zyxw": "专硕", "xsxw": "学硕"}
        degree_text = degree_map.get(item.degree_type, item.degree_type)

        study_mode_map = {"1": "全日制", "2": "非全日制"}
        study_mode_text = study_mode_map.get(item.study_mode, item.study_mode)

        row_dict = {
            "学校代码": item.school_code,
            "学校名称": item.school_name,
            "所在省市": item.province_name,
            "985": "是" if item.is_985 else "否",
            "211": "是" if item.is_211 else "否",
            "双一流": "是" if item.is_double_first_class else "否",
            "自划线": "是" if item.is_self_scoring else "否",
            "博士点": "是" if item.has_doctoral_point else "否",
            "研究生院": "是" if item.has_graduate_school else "否",
            "院系代码": item.college_code,
            "院系名称": item.college_name,
            "专业代码": item.major_code,
            "专业名称": item.major_name,
            "学科门类": item.discipline_category_name,
            "一级学科": item.first_discipline_name,
            "研究方向代码": item.direction_code,
            "研究方向": item.direction_name,
            "学位类型": degree_text,
            "学习方式": study_mode_text,

            "总拟招人数": item.total_planned_enrollment if item.total_planned_enrollment is not None else "",
            "拟招说明": item.planned_enrollment_desc or "",
            "推免人数": item.exempt_student_count if item.exempt_student_count is not None else "",
            "退役士兵计划": "是" if item.is_veteran_plan else "否",
            "少数民族骨干专项计划": "是" if item.is_special_public_plan else "否",

            "科目一(政治/综合)": "\n".join(sub1_list),
            "科目二(外语)": "\n".join(sub2_list),
            "科目三(业务课一)": "\n".join(sub3_list),
            "科目四(业务课二/专业课)": "\n".join(sub4_list),

            "考试方式": item.exam_type_name,
            "指导教师": item.advisor_name,
            "专业备注": item.major_remarks,
        }

        return row_dict

    def _beautify_excel(self, file_path: str) -> None:
        wb = load_workbook(file_path)
        ws = wb.active

        ws.freeze_panes = "A2"

        header_font = Font(name="Microsoft YaHei", size=10, bold=True)
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        left_align_columns = {
            "院系名称", "专业名称", "研究方向", "拟招说明",
            "科目一(政治/综合)", "科目二(外语)", "科目三(业务课一)", "科目四(业务课二/专业课)", "专业备注"
        }

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                col_name = headers[col - 1]
                cell.alignment = align_left if col_name in left_align_columns else align_center

        for col in ws.columns:
            max_width = 0
            for cell in col:
                val = "" if cell.value is None else str(cell.value)
                lines = val.split("\n")
                for line in lines:
                    line_len = sum(2 if ord(char) > 127 else 1 for char in line)
                    if line_len > max_width:
                        max_width = line_len

            col_letter = get_column_letter(col[0].column)
            calculated_width = max_width + 3
            ws.column_dimensions[col_letter].width = min(max(calculated_width, 10), 45)
        wb.save(file_path)

    def run(self) -> None:
        # parse data
        try:
            logger.info("正在解析数据...")
            detail_list = self._get_school_detail_list()
            list_size = len(detail_list)
            logger.info(f"解析成功，共{list_size}条数据")
        except Exception as e:
            logger.critical(f"解析数据失败：{e}")
            sys.exit(1)

        try:
            logger.info("正在写入excel文件...")
            # ensure output path
            os.makedirs(self.output_dir, exist_ok=True)

            # build file name
            time_stamp = time.strftime(r"%Y-%m-%d_%H-%M-%S")
            file_name = f"({self.major_code}){self.major_name}-{time_stamp}.xlsx"
            file_path = os.path.join(self.output_dir, file_name)

            # make excel
            flatten_list = [self._flatten_detail(_) for _ in detail_list]
            df = pd.DataFrame(flatten_list)
            df.to_excel(file_path, index=False, engine="openpyxl")
            self._beautify_excel(file_path)
            logger.info(f"写入完成：{file_path}")
        except Exception as e:
            logger.critical(f"写入excel失败：{e}")
            sys.exit(1)
